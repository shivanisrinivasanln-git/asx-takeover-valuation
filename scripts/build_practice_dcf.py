"""Build the one-sheet practice DCF workbook (docs/practice/one_sheet_dcf.xlsx).

The inputs are the Harbour Industrial example in docs/RESEARCH.md, so the base value per share
matches the hand-worked A$3.05. The sensitivity table is an Excel data table: the cells in it
hold the table formula and Excel fills them in when the workbook is calculated. Excel then
recalculates and saves the file, so the stored values can be tested
(see scripts/recalculate_with_excel.ps1).
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.formula import DataTableFormula
from openpyxl.worksheet.worksheet import Worksheet

NOTE_COLUMN = "H"
PERCENT = "0.0%"
MONEY = "#,##0.0"
YEAR_COLUMNS = ["B", "C", "D", "E", "F"]
TABLE_COLUMNS = ["C", "D", "E", "F", "G"]
GROWTH_VALUES = [0.02, 0.025, 0.03, 0.035, 0.04]
RATE_VALUES = [0.094, 0.099, 0.104, 0.109, 0.114]
FIRST_TABLE_ROW = 27

INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
BOLD = Font(bold=True)


def cell(ws: Worksheet, ref: str) -> Cell:
    target = ws[ref]
    if not isinstance(target, Cell):
        raise TypeError(f"{ref} is not a plain cell")
    return target


def put(
    ws: Worksheet,
    ref: str,
    value: object,
    note: str | None = None,
    number_format: str | None = None,
) -> None:
    target = cell(ws, ref)
    target.value = value  # type: ignore[assignment]
    if number_format:
        target.number_format = number_format
    if note is not None:
        cell(ws, f"{NOTE_COLUMN}{target.row}").value = note


def build(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    if not isinstance(ws, Worksheet):
        raise TypeError("the first sheet is not a worksheet")
    ws.title = "DCF"

    put(ws, "A1", "One-sheet DCF: Harbour Industrial Ltd (invented example from docs/RESEARCH.md)")
    cell(ws, "A1").font = Font(bold=True, size=12)
    put(
        ws,
        "A2",
        "A$ millions unless stated. Yellow cells are inputs. Every formula has a note in column H.",
    )
    put(ws, f"{NOTE_COLUMN}3", "What the formula does")
    cell(ws, f"{NOTE_COLUMN}3").font = BOLD
    put(ws, "A3", "Inputs")
    cell(ws, "A3").font = BOLD

    inputs = [
        (4, "First-year free cash flow", 30.0, MONEY),
        (5, "Growth in free cash flow, years 2 to 5", 0.03, PERCENT),
        (6, "Discount rate (WACC)", 0.104, PERCENT),
        (7, "Terminal growth rate", 0.03, PERCENT),
        (8, "Net debt", 100.0, MONEY),
        (9, "Diluted shares (millions)", 100.0, MONEY),
    ]
    for row, label, value, fmt in inputs:
        put(ws, f"A{row}", label)
        put(ws, f"B{row}", value, "Input", fmt)
        cell(ws, f"B{row}").fill = INPUT_FILL

    put(ws, "A11", "Forecast")
    cell(ws, "A11").font = BOLD
    put(ws, "A12", "Year")
    put(ws, "A13", "Free cash flow")
    put(ws, "A14", "Discount factor")
    put(ws, "A15", "Present value of free cash flow")
    for index, col in enumerate(YEAR_COLUMNS, start=1):
        put(ws, f"{col}12", index)
        cell(ws, f"{col}12").font = BOLD
    put(
        ws,
        f"{NOTE_COLUMN}12",
        "Year numbers 1 to 5; valuation at the start of year 1, cash flows arrive at year end",
    )
    for index, col in enumerate(YEAR_COLUMNS):
        previous = YEAR_COLUMNS[index - 1]
        put(ws, f"{col}13", "=$B$4" if index == 0 else f"={previous}13*(1+$B$5)", None, MONEY)
        put(ws, f"{col}14", f"=1/(1+$B$6)^{col}12", None, "0.0000")
        put(ws, f"{col}15", f"={col}13*{col}14", None, MONEY)
    put(
        ws,
        f"{NOTE_COLUMN}13",
        "Year 1 is the input; each later year grows the year before by the growth input",
    )
    put(
        ws,
        f"{NOTE_COLUMN}14",
        "Discount factor: 1 divided by (1 plus the discount rate) to the power of the year",
    )
    put(ws, f"{NOTE_COLUMN}15", "Present value: free cash flow times the discount factor")

    put(ws, "A17", "Value")
    cell(ws, "A17").font = BOLD
    value_rows = [
        (
            18,
            "Present value of years 1 to 5",
            "=SUM(B15:F15)",
            "Adds the five present values",
            MONEY,
        ),
        (
            19,
            "Terminal value at the end of year 5",
            "=F13*(1+B7)/(B6-B7)",
            "Gordon growth: year 6 free cash flow divided by (discount rate minus terminal growth)",
            MONEY,
        ),
        (
            20,
            "Present value of terminal value",
            "=B19*F14",
            "Discounts the terminal value back with the year 5 discount factor",
            MONEY,
        ),
        (
            21,
            "Enterprise value",
            "=B18+B20",
            "Adds the explicit years and the terminal value",
            MONEY,
        ),
        (22, "Equity value", "=B21-B8", "Enterprise value less net debt", MONEY),
        (23, "Value per share (A$)", "=B22/B9", "Equity value divided by diluted shares", "0.0000"),
    ]
    for row, label, formula, note, fmt in value_rows:
        put(ws, f"A{row}", label)
        put(ws, f"B{row}", formula, note, fmt)
    cell(ws, "A23").font = BOLD
    cell(ws, "B23").font = BOLD

    put(ws, "A25", "Sensitivity: value per share (A$)")
    cell(ws, "A25").font = BOLD
    put(
        ws,
        f"{NOTE_COLUMN}25",
        "Data table: terminal growth across, discount rate down; recalculates B23 for each pair",
    )
    put(ws, "A26", "Discount rate down, terminal growth across")
    put(
        ws,
        "B26",
        "=B23",
        "Corner cell: points at the value per share that the table recalculates",
        "0.0000",
    )
    for col, growth in zip(TABLE_COLUMNS, GROWTH_VALUES, strict=True):
        put(ws, f"{col}26", growth, None, PERCENT)
        cell(ws, f"{col}26").font = BOLD
    for offset, rate in enumerate(RATE_VALUES):
        row = FIRST_TABLE_ROW + offset
        put(ws, f"B{row}", rate, None, PERCENT)
        cell(ws, f"B{row}").font = BOLD
        for col in TABLE_COLUMNS:
            cell(ws, f"{col}{row}").number_format = "0.00"
    last_row = FIRST_TABLE_ROW + len(RATE_VALUES) - 1
    put(
        ws,
        f"C{FIRST_TABLE_ROW}",
        DataTableFormula(
            ref=f"C{FIRST_TABLE_ROW}:G{last_row}", dt2D=True, dtr=True, r1="B7", r2="B6"
        ),
        "Data table C27:G31: row input B7 (terminal growth), column input B6 (discount rate)",
    )

    put(ws, "A33", "Check: value per share by the perpetuity shortcut")
    put(
        ws,
        "B33",
        "=(B4/(B6-B7)-B8)/B9",
        "With equal growth in the forecast and the terminal period, the DCF equals first-year cash "
        "flow over (rate minus growth), less net debt, per share",
        "0.0000",
    )

    ws.column_dimensions["A"].width = 52
    for col in TABLE_COLUMNS:
        ws.column_dimensions[col].width = 12
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions[NOTE_COLUMN].width = 100
    for row_cells in ws.iter_rows(min_row=1, max_row=33):
        for each in row_cells:
            each.alignment = Alignment(vertical="top")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


if __name__ == "__main__":
    build(Path(sys.argv[1]))
