"""The Comparables and Precedents sheets: implied value from a multiple range, as formulas.

Each sheet values Domain at my selected multiple range times FY26 adjusted earnings, then runs
the same equity bridge as the DCF sheet. Below the value block is the evidence: the expert's
printed multiples with their source, and the lowest, median and highest of each column. That
evidence is not used in the value; the value uses the selected range on the Assumptions sheet.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from valuation.inputs import Inputs
from valuation.multiples import Measure, Method

COMPARABLES_SHEET = "Comparables"
PRECEDENTS_SHEET = "Precedents"
NOTE_COLUMN = "J"
MONEY = "#,##0.0"
MULTIPLE = "0.0"
PERCENT = "0.00"
PER_SHARE = "0.00"
BOLD = Font(bold=True)
HEADING_FILL = PatternFill("solid", start_color="D9E1F2")
CASE_COLUMNS = {"low": "C", "mid": "D", "high": "E"}
MEASURES: tuple[Measure, ...] = ("ebitda", "ebita")
FIRST_BLOCK_ROW = 4
BLOCK_HEIGHT = 14
# The four evidence columns: (measure, basis) in the order they are written.
EVIDENCE_COLUMNS = (
    ("ebitda", "historical"),
    ("ebitda", "forecast"),
    ("ebita", "historical"),
    ("ebita", "forecast"),
)
EVIDENCE_HEADERS = (
    "EBITDA, historical",
    "EBITDA, forecast",
    "EBITA, historical",
    "EBITA, forecast",
)

# Offsets of the rows inside one measure block.
HEAD, EARNINGS, MULTIPLE_ROW, BEFORE, PREMIUM, OPERATIONS = 0, 1, 2, 3, 4, 5
OTHER, EV, NET_DEBT, EQUITY, SHARES, PER_SHARE_ROW = 6, 7, 8, 9, 10, 11


def _cell(sheet: Worksheet, ref: str) -> Cell:
    found = sheet[ref]
    if not isinstance(found, Cell):
        raise TypeError(f"{ref} is not a plain cell")
    return found


def _put(sheet: Worksheet, ref: str, value: object, number_format: str | None = None) -> None:
    cell = _cell(sheet, ref)
    cell.value = value  # type: ignore[assignment]
    if number_format:
        cell.number_format = number_format


def _heading(sheet: Worksheet, row: int, text: str, last_column: int) -> None:
    sheet[f"A{row}"] = text
    for column in range(1, last_column + 1):
        cell = sheet.cell(row, column)
        cell.font = BOLD
        cell.fill = HEADING_FILL


def _value_block(
    sheet: Worksheet, method: Method, measure: Measure, first_row: int
) -> dict[str, str]:
    """One measure's value block. Returns the cell of each output, keyed by valuation.json key."""
    prefix = f"{method}_{measure}"
    row = {
        name: first_row + offset
        for name, offset in (
            ("head", HEAD),
            ("earnings", EARNINGS),
            ("multiple", MULTIPLE_ROW),
            ("before", BEFORE),
            ("premium", PREMIUM),
            ("operations", OPERATIONS),
            ("other", OTHER),
            ("ev", EV),
            ("net_debt", NET_DEBT),
            ("equity", EQUITY),
            ("shares", SHARES),
            ("per_share", PER_SHARE_ROW),
        )
    }
    _heading(sheet, row["head"], f"{measure.upper()} multiple (A$ millions unless stated)", 5)
    for case, column in CASE_COLUMNS.items():
        _put(sheet, f"{column}{row['head']}", f"{case.capitalize()} case")
        sheet[f"{column}{row['head']}"].font = BOLD

    labels = {
        "earnings": (f"FY26 adjusted {measure.upper()}", "A$m"),
        "multiple": ("Selected multiple", "x"),
        "before": ("Value before any control premium", "A$m"),
        "premium": ("Control premium applied", "%"),
        "operations": ("Value of business operations", "A$m"),
        "other": ("Other assets and liabilities", "A$m"),
        "ev": ("Enterprise value", "A$m"),
        "net_debt": ("Adjusted net debt", "A$m"),
        "equity": ("Equity value", "A$m"),
        "shares": ("Diluted shares", "m"),
        "per_share": ("Value per share", "A$"),
    }
    for name, (label, unit) in labels.items():
        sheet[f"A{row[name]}"] = label
        sheet[f"B{row[name]}"] = unit

    if method == "trading":
        premium_note = (
            "My control premium input, applied here only because peer multiples come from "
            "minority share prices."
        )
    else:
        premium_note = "Deal multiples are control multiples, so no premium is applied."
    notes = {
        "earnings": "The expert's FY26 adjusted earnings, after the dividends to agents.",
        "multiple": (
            f"The low and high ends of my selected {method} range; the mid case is their midpoint."
        ),
        "before": "Earnings times the multiple.",
        "premium": premium_note,
        "operations": "Value before premium grossed up by the premium.",
        "other": "Equity accounted investments, provisions and contingent consideration.",
        "ev": "Operations value plus other assets and liabilities.",
        "net_debt": "The same adjusted net debt as the DCF sheet.",
        "equity": "Enterprise value less adjusted net debt.",
        "shares": "Larger count in the low case, the average in the mid case, as on the DCF sheet.",
        "per_share": "Equity value over diluted shares.",
    }
    for name, text in notes.items():
        sheet[f"{NOTE_COLUMN}{row[name]}"] = text

    outputs: dict[str, str] = {}
    for case, column in CASE_COLUMNS.items():
        r = {name: f"{column}{number}" for name, number in row.items()}
        _put(sheet, r["earnings"], f"=fy26_adjusted_{measure}_aud_m", MONEY)
        if case == "mid":
            low_cell = f"{CASE_COLUMNS['low']}{row['multiple']}"
            high_cell = f"{CASE_COLUMNS['high']}{row['multiple']}"
            _put(sheet, r["multiple"], f"=({low_cell}+{high_cell})/2", MULTIPLE)
        else:
            _put(sheet, r["multiple"], f"={method}_{measure}_multiple_fy26_{case}_x", MULTIPLE)
        _put(sheet, r["before"], f"={r['earnings']}*{r['multiple']}", MONEY)
        if method == "trading":
            _put(sheet, r["premium"], "=control_premium_applied_pct", PERCENT)
        else:
            _put(sheet, r["premium"], 0, PERCENT)
        _put(sheet, r["operations"], f"={r['before']}*(1+{r['premium']}/100)", MONEY)
        _put(sheet, r["other"], "=other_assets_and_liabilities_aud_m", MONEY)
        _put(sheet, r["ev"], f"={r['operations']}+{r['other']}", MONEY)
        _put(sheet, r["net_debt"], "=adjusted_net_debt_aud_m", MONEY)
        _put(sheet, r["equity"], f"={r['ev']}-{r['net_debt']}", MONEY)
        _put(sheet, r["shares"], f"=dcf_diluted_shares_m_{case}", "#,##0.000")
        _put(sheet, r["per_share"], f"={r['equity']}/{r['shares']}", PER_SHARE)
        stems = {
            f"{prefix}_multiple_x_{case}": "multiple",
            f"{prefix}_business_operations_value_aud_m_{case}": "operations",
            f"{prefix}_enterprise_value_aud_m_{case}": "ev",
            f"{prefix}_equity_value_aud_m_{case}": "equity",
            f"{prefix}_diluted_shares_m_{case}": "shares",
            f"{prefix}_value_per_share_aud_{case}": "per_share",
        }
        for key, name in stems.items():
            outputs[key] = _absolute(sheet, f"{column}{row[name]}")
    outputs[f"{prefix}_earnings_aud_m"] = _absolute(sheet, f"C{row['earnings']}")
    outputs[f"{prefix}_control_premium_pct"] = _absolute(sheet, f"C{row['premium']}")
    for column_number in range(1, 6):
        sheet.cell(row["per_share"], column_number).font = BOLD
    return outputs


def _absolute(sheet: Worksheet, ref: str) -> str:
    column = "".join(ch for ch in ref if ch.isalpha())
    number = ref[len(column) :]
    return f"{sheet.title}!${column}${number}"


def _evidence_table(
    sheet: Worksheet,
    first_row: int,
    title: str,
    first_header: str,
    rows: Sequence[Mapping[str, str]],
    group_key: str,
) -> None:
    """Pivot the transcribed (name, measure, basis, multiple) rows into one row per name."""
    _heading(sheet, first_row, title, 9)
    header = first_row + 1
    sheet[f"A{header}"] = first_header
    for offset, text in enumerate(EVIDENCE_HEADERS):
        sheet.cell(header, 3 + offset, text).font = BOLD
    sheet[f"G{header}"] = "Source"
    sheet[f"H{header}"] = "Location"
    sheet[f"I{header}"] = "As at"
    for column in "AG":
        sheet[f"{column}{header}"].font = BOLD
    sheet[f"H{header}"].font = BOLD
    sheet[f"I{header}"].font = BOLD

    names: list[str] = []
    for row in rows:
        if row[group_key] not in names:
            names.append(row[group_key])
    for offset, name in enumerate(names):
        number = header + 1 + offset
        mine = [row for row in rows if row[group_key] == name]
        sheet[f"A{number}"] = name
        for column_offset, (measure, basis) in enumerate(EVIDENCE_COLUMNS):
            for row in mine:
                if row["measure"] == measure and row["basis"] == basis:
                    sheet.cell(
                        number, 3 + column_offset, float(row["multiple_x"])
                    ).number_format = MULTIPLE
        sheet[f"G{number}"] = mine[0]["source"]
        sheet[f"H{number}"] = "; ".join(dict.fromkeys(row["location"] for row in mine))
        sheet[f"I{number}"] = mine[0]["as_at"]
    last = header + len(names)
    for offset, (label, function) in enumerate(
        (("Lowest", "MIN"), ("Median", "MEDIAN"), ("Highest", "MAX"))
    ):
        number = last + 1 + offset
        sheet[f"A{number}"] = label
        sheet[f"A{number}"].font = BOLD
        for column_offset in range(len(EVIDENCE_COLUMNS)):
            letter = chr(ord("C") + column_offset)
            _put(
                sheet,
                f"{letter}{number}",
                f"={function}({letter}{header + 1}:{letter}{last})",
                MULTIPLE,
            )
        sheet[f"{NOTE_COLUMN}{number}"] = (
            f"{label} of the column. Evidence only: the value above uses my selected range."
        )


def _setup(sheet: Worksheet, title: str, description: str) -> None:
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = description
    sheet.column_dimensions["A"].width = 44
    sheet.column_dimensions["B"].width = 9
    for column in "CDEF":
        sheet.column_dimensions[column].width = 18
    sheet.column_dimensions["G"].width = 40
    sheet.column_dimensions["H"].width = 34
    sheet.column_dimensions["I"].width = 14
    sheet.column_dimensions[NOTE_COLUMN].width = 90


def add_multiples_sheet(book: Workbook, inputs: Inputs, method: Method) -> dict[str, str]:
    """Add the Comparables or Precedents sheet; returns each output cell by valuation.json key."""
    if method == "trading":
        sheet: Worksheet = book.create_sheet(COMPARABLES_SHEET)
        _setup(
            sheet,
            "Trading comparables",
            "Domain valued at my selected peer multiple range times FY26 adjusted earnings. "
            "Peer multiples come from share prices at 30 May 2025 and carry no control premium.",
        )
    else:
        sheet = book.create_sheet(PRECEDENTS_SHEET)
        _setup(
            sheet,
            "Precedent transactions",
            "Domain valued at my selected deal multiple range times FY26 adjusted earnings. "
            "Deal multiples include a control premium.",
        )
    outputs: dict[str, str] = {}
    for index, measure in enumerate(MEASURES):
        outputs.update(
            _value_block(sheet, method, measure, FIRST_BLOCK_ROW + index * (BLOCK_HEIGHT))
        )
    evidence_row = FIRST_BLOCK_ROW + len(MEASURES) * BLOCK_HEIGHT + 1
    if method == "trading":
        _evidence_table(
            sheet,
            evidence_row,
            "The expert's peer multiples at 30 May 2025 (evidence)",
            "Company",
            inputs.tables["comparable_multiples"],
            "company",
        )
    else:
        _evidence_table(
            sheet,
            evidence_row,
            "The expert's deal multiples (evidence)",
            "Deal",
            _deal_rows(inputs.tables["precedent_transactions"]),
            "deal",
        )
    return outputs


def _deal_rows(rows: Sequence[Mapping[str, str]]) -> list[Mapping[str, str]]:
    """Name each deal with its year so two deals with the same name stay apart."""
    return [{**row, "deal": f"{row['deal']} ({row['announced']})"} for row in rows]
