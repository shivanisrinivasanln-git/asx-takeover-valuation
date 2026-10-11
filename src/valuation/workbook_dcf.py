"""The DCF sheet: the free cash flow forecast, the value at three discount rates and the
sensitivity grid, all as formulas that read the Assumptions sheet through defined names.

The formulas follow `forecast.py`, `dcf.py`, `bridge.py` and `sensitivity.py` line by line, so a
reader can hold the Python and the sheet side by side. The only numbers typed into a formula are
0, 1, 2 and 100 (a unit conversion, a step of one year or an average of two), which a test
enforces. The sheet returns the cell of every output, keyed by its `valuation.json` key.
"""

from __future__ import annotations

from collections.abc import Callable

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

DCF_SHEET = "DCF"
NOTE_COLUMN = "N"
LABEL_WIDTH = 50
YEARS = 11  # the stub period (year 0) and FY26 to FY35
FIRST_YEAR_COLUMN = 3
MONEY = "#,##0.0"
PERCENT = "0.00"
FACTOR = "0.0000"
PER_SHARE = "0.00"
BOLD = Font(bold=True)
HEADING_FILL = PatternFill("solid", start_color="D9E1F2")
CASES = ("low", "mid", "high")

# Row numbers of the forecast block.
R_YEARS = 3
R_YEAR = 5
R_PERIOD = 6
R_FRACTION = 7
R_FULL_REVENUE = 8
R_PRIOR_REVENUE = 9
R_REVENUE = 10
R_MARGIN = 11
R_EBITDA = 12
R_SAVING = 13
R_DEP_PCT = 14
R_DEP = 15
R_TAXABLE = 16
R_TAX = 17
R_CAPEX_PCT = 18
R_CAPEX = 19
R_WC = 20
R_AGENT = 21
R_FCF = 22
R_TIME = 23

# Value block: columns C, D and E hold the low, mid and high cases.
R_VALUE_HEAD = 26
R_RATE = 28
R_GROWTH = 29
R_PV_FLOWS = 30
R_TV = 31
R_TV_FACTOR = 32
R_PV_TV = 33
R_OPERATIONS = 34
R_TV_SHARE = 35
R_OTHER = 36
R_EV = 37
R_NET_DEBT = 38
R_EQUITY = 39
R_SHARES = 40
R_PER_SHARE = 41

# Terminal year block (one column, at the terminal growth rate).
R_TERMINAL_HEAD = 44
TERMINAL_FIRST_ROW = 45

# Sensitivity block.
R_SENS_HEAD = 58
R_SENS_OFFSET = 59
R_SENS_GROWTH = 60
SENS_FIRST_ROW = 61
SENS_GRID_FIRST_ROW = 73
SENS_SIZE = 5
SENS_PV_COLUMN = "H"

# Discount rate build: one block for my risk-free rate and one for the expert's.
R_WACC_HEAD = 80
WACC_BLOCK_HEIGHT = 10


def _first_year_column() -> str:
    return get_column_letter(FIRST_YEAR_COLUMN)


def _year_column(year: int) -> str:
    return get_column_letter(FIRST_YEAR_COLUMN + year)


def _last_year_column() -> str:
    return _year_column(YEARS - 1)


def _row_range(row: int) -> str:
    return f"${_first_year_column()}${row}:${_last_year_column()}${row}"


def _cell(sheet: Worksheet, ref: str) -> Cell:
    found = sheet[ref]
    if not isinstance(found, Cell):
        raise TypeError(f"{ref} is not a plain cell")
    return found


def _put(sheet: Worksheet, ref: str, value: object, number_format: str | None = None) -> None:
    cell = _cell(sheet, ref)
    cell.value = value  # type: ignore[assignment]
    if number_format:
        cell.number_format = number_format


def _note(sheet: Worksheet, row: int, text: str) -> None:
    sheet[f"{NOTE_COLUMN}{row}"] = text


def _heading(sheet: Worksheet, row: int, text: str) -> None:
    sheet[f"A{row}"] = text
    for column in range(1, FIRST_YEAR_COLUMN + YEARS):
        cell = sheet.cell(row, column)
        cell.font = BOLD
        cell.fill = HEADING_FILL


def _year_row(
    sheet: Worksheet,
    row: int,
    label: str,
    unit: str,
    stub: str,
    later: Callable[[str, str], str],
    note: str,
    number_format: str,
) -> None:
    """A row with one formula for the stub column and one pattern for FY26 to FY35.

    `later` receives the column letter and the letter of the column before it.
    """
    sheet[f"A{row}"] = label
    sheet[f"B{row}"] = unit
    _put(sheet, f"{_year_column(0)}{row}", stub, number_format)
    for year in range(1, YEARS):
        column, previous = _year_column(year), _year_column(year - 1)
        _put(sheet, f"{column}{row}", later(column, previous), number_format)
    _note(sheet, row, note)


def _forecast_block(sheet: Worksheet) -> None:
    first = _first_year_column()
    _heading(sheet, R_YEAR - 1, "Free cash flow forecast (A$ millions unless stated)")
    sheet[f"A{R_YEARS}"] = "Full forecast years after the stub period"
    sheet[f"B{R_YEARS}"] = "years"
    _put(
        sheet,
        f"{first}{R_YEARS}",
        "=INDEX(projection_end_date_ymd,1)-(INDEX(valuation_date_ymd,1)+1)",
        "0",
    )
    _note(sheet, R_YEARS, "Year of the projection end date less the year after the valuation date.")

    sheet[f"A{R_YEAR}"] = "Year after the stub period"
    sheet[f"B{R_YEAR}"] = "number"
    sheet[f"{first}{R_YEAR}"] = 0
    for year in range(1, YEARS):
        sheet[f"{_year_column(year)}{R_YEAR}"] = f"={_year_column(year - 1)}{R_YEAR}+1"
    _note(sheet, R_YEAR, "Year 0 is the second half of FY25, which the cash flows start from.")

    _year_row(
        sheet,
        R_PERIOD,
        "Period",
        "label",
        f'="FY"&(INDEX(valuation_date_ymd,1)+1+{first}{R_YEAR})&" second half"',
        lambda c, _p: f'="FY"&(INDEX(valuation_date_ymd,1)+1+{c}{R_YEAR})',
        "Fiscal year label, from the valuation date.",
        "General",
    )
    _year_row(
        sheet,
        R_FRACTION,
        "Fraction of a year in the period",
        "years",
        "=stub_period_years",
        lambda c, _p: f"=IF({c}{R_YEAR}=0,stub_period_years,1)",
        "The stub period input for the first period; every later period is a full year.",
        "0.00",
    )
    _year_row(
        sheet,
        R_FULL_REVENUE,
        "Full-year revenue",
        "A$m",
        "=fy25_revenue_aud_m",
        lambda c, p: f"={p}{R_FULL_REVENUE}*(1+revenue_growth_fy26_to_fy35_pct/100)",
        "FY25 is the broker median; each later year grows at the revenue growth input.",
        MONEY,
    )
    _year_row(
        sheet,
        R_PRIOR_REVENUE,
        "Prior full-year revenue",
        "A$m",
        "=fy24_revenue_aud_m",
        lambda c, p: f"={p}{R_FULL_REVENUE}",
        "The year before; used for the increase in working capital and the agents' dividends.",
        MONEY,
    )
    _year_row(
        sheet,
        R_REVENUE,
        "Revenue in the period",
        "A$m",
        f"={first}{R_FULL_REVENUE}-first_half_fy25_revenue_aud_m",
        lambda c, _p: f"={c}{R_FULL_REVENUE}",
        "Stub revenue is the FY25 forecast less the first half already reported.",
        MONEY,
    )
    _year_row(
        sheet,
        R_MARGIN,
        "EBITDA margin (full year)",
        "%",
        "=fy25_ebitda_aud_m/fy25_revenue_aud_m*100",
        lambda c, _p: (
            f"=IF({c}{R_YEAR}<=early_years_count,"
            f"${first}{R_MARGIN}+(ebitda_margin_fy30_pct-${first}{R_MARGIN})"
            f"*{c}{R_YEAR}/early_years_count,"
            f"ebitda_margin_fy30_pct+(ebitda_margin_fy35_pct-ebitda_margin_fy30_pct)"
            f"*({c}{R_YEAR}-early_years_count)/(${first}${R_YEARS}-early_years_count))"
        ),
        "Straight line from the FY25 margin to the FY30 margin, then to the FY35 margin.",
        PERCENT,
    )
    _year_row(
        sheet,
        R_EBITDA,
        "EBITDA",
        "A$m",
        "=fy25_ebitda_aud_m-first_half_fy25_ebitda_aud_m",
        lambda c, _p: f"={c}{R_REVENUE}*{c}{R_MARGIN}/100",
        "Stub EBITDA is the FY25 forecast less the first half; later it is revenue times margin.",
        MONEY,
    )
    _year_row(
        sheet,
        R_SAVING,
        "Corporate cost saving",
        "A$m",
        f"=corporate_cost_saving_aud_m*{first}{R_FRACTION}",
        lambda c, _p: f"=corporate_cost_saving_aud_m*{c}{R_FRACTION}",
        "The annual saving input, scaled by the part of the year in the period.",
        MONEY,
    )
    _year_row(
        sheet,
        R_DEP_PCT,
        "Depreciation and amortisation (full year, % of revenue)",
        "%",
        "=fy25_depreciation_and_other_amortisation_aud_m/fy25_revenue_aud_m*100",
        lambda c, _p: (
            f"=IF({c}{R_YEAR}<=early_years_count,"
            f"${first}{R_DEP_PCT}+(depreciation_pct_of_revenue_long_term_pct-${first}{R_DEP_PCT})"
            f"*{c}{R_YEAR}/early_years_count,depreciation_pct_of_revenue_long_term_pct)"
        ),
        "Straight line from the FY25 share to the long-term share, reached after the early years.",
        PERCENT,
    )
    _year_row(
        sheet,
        R_DEP,
        "Depreciation and amortisation",
        "A$m",
        "=fy25_depreciation_and_other_amortisation_aud_m"
        "+first_half_fy25_depreciation_and_other_amortisation_aud_m",
        lambda c, _p: f"={c}{R_REVENUE}*{c}{R_DEP_PCT}/100",
        "The first-half figure is stored as printed, a negative cost, so adding it removes it.",
        MONEY,
    )
    _year_row(
        sheet,
        R_TAXABLE,
        "Taxable profit",
        "A$m",
        f"={first}{R_EBITDA}+{first}{R_SAVING}-{first}{R_DEP}",
        lambda c, _p: f"={c}{R_EBITDA}+{c}{R_SAVING}-{c}{R_DEP}",
        "EBITDA plus the saving, less depreciation; no deduction for acquired intangibles.",
        MONEY,
    )
    _year_row(
        sheet,
        R_TAX,
        "Tax",
        "A$m",
        f"={first}{R_TAXABLE}*tax_rate_pct/100",
        lambda c, _p: f"={c}{R_TAXABLE}*tax_rate_pct/100",
        "Taxable profit times the tax rate.",
        MONEY,
    )
    _year_row(
        sheet,
        R_CAPEX_PCT,
        "Capital expenditure (% of revenue)",
        "%",
        "=capex_pct_of_revenue_fy26_to_fy30_pct",
        lambda c, _p: (
            f"=IF({c}{R_YEAR}<=early_years_count,capex_pct_of_revenue_fy26_to_fy30_pct,"
            "capex_pct_of_revenue_long_term_pct)"
        ),
        "The early-year share for the stub period and FY26 to FY30, then the long-term share.",
        PERCENT,
    )
    _year_row(
        sheet,
        R_CAPEX,
        "Capital expenditure",
        "A$m",
        f"={first}{R_REVENUE}*{first}{R_CAPEX_PCT}/100",
        lambda c, _p: f"={c}{R_REVENUE}*{c}{R_CAPEX_PCT}/100",
        "Revenue in the period times the capital expenditure share.",
        MONEY,
    )
    _year_row(
        sheet,
        R_WC,
        "Increase in working capital",
        "A$m",
        f"=({first}{R_FULL_REVENUE}-{first}{R_PRIOR_REVENUE})*working_capital_pct_of_revenue_pct"
        f"/100*{first}{R_FRACTION}",
        lambda c, _p: (
            f"=({c}{R_FULL_REVENUE}-{c}{R_PRIOR_REVENUE})*working_capital_pct_of_revenue_pct"
            f"/100*{c}{R_FRACTION}"
        ),
        "The working capital share times the change in full-year revenue, scaled for the period.",
        MONEY,
    )
    _year_row(
        sheet,
        R_AGENT,
        "Dividends to agents",
        "A$m",
        f"={first}{R_PRIOR_REVENUE}*residential_and_mdc_share_of_revenue_pct/100"
        f"*agent_dividend_pct_of_prior_year_revenue_pct/100*{first}{R_FRACTION}",
        lambda c, _p: (
            f"={c}{R_PRIOR_REVENUE}*residential_and_mdc_share_of_revenue_pct/100"
            f"*agent_dividend_pct_of_prior_year_revenue_pct/100*{c}{R_FRACTION}"
        ),
        "Paid on the prior year's revenue of the segments the agents hold equity in.",
        MONEY,
    )
    _year_row(
        sheet,
        R_FCF,
        "Free cash flow",
        "A$m",
        f"={first}{R_TAXABLE}-{first}{R_TAX}+{first}{R_DEP}-{first}{R_CAPEX}-{first}{R_WC}"
        f"-{first}{R_AGENT}",
        lambda c, _p: f"={c}{R_TAXABLE}-{c}{R_TAX}+{c}{R_DEP}-{c}{R_CAPEX}-{c}{R_WC}-{c}{R_AGENT}",
        "Profit after tax plus depreciation, less capital expenditure, working capital, dividends.",
        MONEY,
    )
    for column in range(1, FIRST_YEAR_COLUMN + YEARS):
        sheet.cell(R_FCF, column).font = BOLD
    _year_row(
        sheet,
        R_TIME,
        "Years from the valuation date to the cash flow",
        "years",
        "=stub_period_years*(1-mid_year_discounting_flag/2)",
        lambda c, _p: f"=stub_period_years+{c}{R_YEAR}-mid_year_discounting_flag/2",
        "Mid-year timing counts each flow from the middle of its period; year end, from the end.",
        "0.00",
    )


def _terminal_year(
    sheet: Worksheet, column: str, first_row: int, growth_ref: str
) -> dict[str, int]:
    """Write the free cash flow of the year after FY35 at a terminal growth rate. Returns rows."""
    last = _last_year_column()
    rows = {
        "revenue": first_row,
        "ebitda": first_row + 1,
        "saving": first_row + 2,
        "dep": first_row + 3,
        "taxable": first_row + 4,
        "tax": first_row + 5,
        "capex": first_row + 6,
        "wc": first_row + 7,
        "agent": first_row + 8,
        "fcf": first_row + 9,
    }
    r = rows
    formulas = {
        "revenue": f"=${last}${R_REVENUE}*(1+{growth_ref}/100)",
        "ebitda": f"={column}{r['revenue']}*${last}${R_EBITDA}/${last}${R_REVENUE}",
        "saving": f"=${last}${R_SAVING}",
        "dep": f"={column}{r['revenue']}*depreciation_pct_of_revenue_long_term_pct/100",
        "taxable": f"={column}{r['ebitda']}+{column}{r['saving']}-{column}{r['dep']}",
        "tax": f"={column}{r['taxable']}*tax_rate_pct/100",
        "capex": f"={column}{r['revenue']}*capex_pct_of_revenue_long_term_pct/100",
        "wc": (
            f"=({column}{r['revenue']}-${last}${R_REVENUE})*working_capital_pct_of_revenue_pct/100"
        ),
        "agent": (
            f"=${last}${R_REVENUE}*residential_and_mdc_share_of_revenue_pct/100"
            "*agent_dividend_pct_of_prior_year_revenue_pct/100"
        ),
        "fcf": (
            f"={column}{r['taxable']}-{column}{r['tax']}+{column}{r['dep']}"
            f"-{column}{r['capex']}-{column}{r['wc']}-{column}{r['agent']}"
        ),
    }
    for name, row in rows.items():
        _put(sheet, f"{column}{row}", formulas[name], MONEY)
    return rows


TERMINAL_LABELS = (
    ("Revenue", "A$m", "Last forecast year's revenue grown at the terminal growth rate."),
    ("EBITDA", "A$m", "The last forecast year's EBITDA margin is held."),
    ("Corporate cost saving", "A$m", "The last forecast year's saving."),
    ("Depreciation and amortisation", "A$m", "The long-term share of revenue."),
    ("Taxable profit", "A$m", "EBITDA plus the saving, less depreciation."),
    ("Tax", "A$m", "Taxable profit times the tax rate."),
    ("Capital expenditure", "A$m", "The long-term share of revenue."),
    (
        "Increase in working capital",
        "A$m",
        "The working capital share of the smaller revenue increase.",
    ),
    ("Dividends to agents", "A$m", "Paid on the last forecast year's revenue."),
    (
        "Free cash flow in the year after the forecast",
        "A$m",
        "The cash flow the terminal value grows from.",
    ),
)


def _terminal_labels(sheet: Worksheet, first_row: int) -> None:
    for offset, (label, unit, note) in enumerate(TERMINAL_LABELS):
        sheet[f"A{first_row + offset}"] = label
        sheet[f"B{first_row + offset}"] = unit
        _note(sheet, first_row + offset, note)


def _value_block(sheet: Worksheet, terminal_fcf_ref: str) -> dict[str, dict[str, str]]:
    """The value at the low, mid and high discount rates. Returns the cell of each output."""
    _heading(sheet, R_VALUE_HEAD, "Value at three discount rates (A$ millions unless stated)")
    for case, column in zip(CASES, "CDE", strict=True):
        sheet[f"{column}{R_VALUE_HEAD + 1}"] = f"{case.capitalize()} case"
        sheet[f"{column}{R_VALUE_HEAD + 1}"].font = BOLD
    _note(
        sheet,
        R_VALUE_HEAD + 1,
        "The low case discounts at the top of the range and uses more shares.",
    )

    time_range = _row_range(R_TIME)
    fcf_range = _row_range(R_FCF)
    last = _last_year_column()
    labels = {
        R_RATE: ("Discount rate", "%"),
        R_GROWTH: ("Terminal growth rate", "%"),
        R_PV_FLOWS: ("Present value of the cash flows", "A$m"),
        R_TV: ("Terminal value at the end of the forecast", "A$m"),
        R_TV_FACTOR: ("Discount factor for the terminal value", "factor"),
        R_PV_TV: ("Present value of the terminal value", "A$m"),
        R_OPERATIONS: ("Value of business operations", "A$m"),
        R_TV_SHARE: ("Terminal value share of operations value", "%"),
        R_OTHER: ("Other assets and liabilities", "A$m"),
        R_EV: ("Enterprise value", "A$m"),
        R_NET_DEBT: ("Adjusted net debt", "A$m"),
        R_EQUITY: ("Equity value", "A$m"),
        R_SHARES: ("Diluted shares", "m"),
        R_PER_SHARE: ("Value per share", "A$"),
    }
    for row, (label, unit) in labels.items():
        sheet[f"A{row}"] = label
        sheet[f"B{row}"] = unit
    notes = {
        R_RATE: "Low: top of the adopted range. Mid: its middle. High: bottom of the range.",
        R_GROWTH: "The terminal growth input, the same in all three cases.",
        R_PV_FLOWS: "Each free cash flow divided by one plus the rate, to the power of its time.",
        R_TV: "Gordon growth: the next flow over (rate less growth), raised for mid-year timing.",
        R_TV_FACTOR: "Discounted from the end of the forecast: the stub plus the full years.",
        R_PV_TV: "Terminal value times its discount factor.",
        R_OPERATIONS: "Present value of the cash flows plus that of the terminal value.",
        R_TV_SHARE: "How much of the value comes after the forecast.",
        R_OTHER: "Equity accounted investments, provisions and contingent consideration.",
        R_EV: "Operations value plus other assets and liabilities.",
        R_NET_DEBT: "Reported net debt including leases, plus the expert's two adjustments.",
        R_EQUITY: "Enterprise value less adjusted net debt.",
        R_SHARES: "Ordinary and cash-settled rights less treasury shares, plus rights that vest.",
        R_PER_SHARE: "Equity value over diluted shares.",
    }
    for row, text in notes.items():
        _note(sheet, row, text)

    discount_rates = {
        "C": "=discount_rate_high_pct",
        "D": "=(discount_rate_low_pct+discount_rate_high_pct)/2",
        "E": "=discount_rate_low_pct",
    }
    fixed_shares = "ordinary_shares_on_issue_m+share_rights_cash_settled_m-treasury_shares_m"
    shares = {
        "C": f"={fixed_shares}+share_rights_vesting_low_m",
        "D": f"=(C{R_SHARES}+E{R_SHARES})/2",
        "E": f"={fixed_shares}+share_rights_vesting_high_m",
    }
    for column in "CDE":
        rate = f"{column}{R_RATE}"
        growth = f"{column}{R_GROWTH}"
        _put(sheet, rate, discount_rates[column], PERCENT)
        _put(sheet, growth, "=terminal_growth_pct", PERCENT)
        _put(
            sheet,
            f"{column}{R_PV_FLOWS}",
            f"=SUMPRODUCT({fcf_range}/(1+{rate}/100)^{time_range})",
            MONEY,
        )
        _put(
            sheet,
            f"{column}{R_TV}",
            f"={terminal_fcf_ref}/(({rate}-{growth})/100)*(1+{rate}/100)^(mid_year_discounting_flag/2)",
            MONEY,
        )
        _put(
            sheet,
            f"{column}{R_TV_FACTOR}",
            f"=1/(1+{rate}/100)^(stub_period_years+${last}${R_YEAR})",
            FACTOR,
        )
        _put(sheet, f"{column}{R_PV_TV}", f"={column}{R_TV}*{column}{R_TV_FACTOR}", MONEY)
        _put(sheet, f"{column}{R_OPERATIONS}", f"={column}{R_PV_FLOWS}+{column}{R_PV_TV}", MONEY)
        _put(
            sheet,
            f"{column}{R_TV_SHARE}",
            f"={column}{R_PV_TV}/{column}{R_OPERATIONS}*100",
            PERCENT,
        )
        _put(sheet, f"{column}{R_OTHER}", "=other_assets_and_liabilities_aud_m", MONEY)
        _put(sheet, f"{column}{R_EV}", f"={column}{R_OPERATIONS}+{column}{R_OTHER}", MONEY)
        _put(
            sheet,
            f"{column}{R_NET_DEBT}",
            "=reported_net_debt_including_leases_aud_m"
            "+net_debt_adjustment_capitalised_borrowing_costs_aud_m"
            "+net_debt_adjustment_interim_dividend_aud_m",
            MONEY,
        )
        _put(sheet, f"{column}{R_EQUITY}", f"={column}{R_EV}-{column}{R_NET_DEBT}", MONEY)
        _put(sheet, f"{column}{R_SHARES}", shares[column], "#,##0.000")
        _put(sheet, f"{column}{R_PER_SHARE}", f"={column}{R_EQUITY}/{column}{R_SHARES}", PER_SHARE)
    for column_number in range(1, FIRST_YEAR_COLUMN + YEARS):
        sheet.cell(R_PER_SHARE, column_number).font = BOLD

    stems = {
        "dcf_discount_rate_pct": R_RATE,
        "dcf_present_value_of_cash_flows_aud_m": R_PV_FLOWS,
        "dcf_terminal_value_aud_m": R_TV,
        "dcf_present_value_of_terminal_value_aud_m": R_PV_TV,
        "dcf_terminal_value_share_of_operations_pct": R_TV_SHARE,
        "dcf_business_operations_value_aud_m": R_OPERATIONS,
        "dcf_enterprise_value_aud_m": R_EV,
        "dcf_equity_value_aud_m": R_EQUITY,
        "dcf_diluted_shares_m": R_SHARES,
        "dcf_value_per_share_aud": R_PER_SHARE,
    }
    return {
        case: {f"{stem}_{case}": f"{column}{row}" for stem, row in stems.items()}
        for case, column in zip(CASES, "CDE", strict=True)
    }


def _sensitivity_block(sheet: Worksheet) -> None:
    last = _last_year_column()
    _heading(sheet, R_SENS_HEAD, "Sensitivity of the value per share to the rate and growth")
    sheet[f"A{R_SENS_OFFSET}"] = "Steps from the base case"
    sheet[f"B{R_SENS_OFFSET}"] = "number"
    sheet[f"A{R_SENS_GROWTH}"] = "Terminal growth rate"
    sheet[f"B{R_SENS_GROWTH}"] = "%"
    columns = [get_column_letter(FIRST_YEAR_COLUMN + i) for i in range(SENS_SIZE)]
    for index, column in enumerate(columns):
        offset = (
            "=-sensitivity_steps_each_side"
            if index == 0
            else f"={columns[index - 1]}{R_SENS_OFFSET}+1"
        )
        _put(sheet, f"{column}{R_SENS_OFFSET}", offset, "0")
        _put(
            sheet,
            f"{column}{R_SENS_GROWTH}",
            f"=terminal_growth_pct+{column}{R_SENS_OFFSET}*sensitivity_growth_step_pct",
            PERCENT,
        )
        _terminal_year(sheet, column, SENS_FIRST_ROW, f"{column}${R_SENS_GROWTH}")
    _note(sheet, R_SENS_OFFSET, "From minus the steps each side to plus the steps each side.")
    _note(sheet, R_SENS_GROWTH, "The terminal growth input, moved by the growth step.")
    _terminal_labels(sheet, SENS_FIRST_ROW)
    for line, (_, _, text) in enumerate(TERMINAL_LABELS):
        _note(sheet, SENS_FIRST_ROW + line, text + " Repeated for each growth rate in the grid.")

    grid_head = SENS_GRID_FIRST_ROW - 1
    sheet[f"A{grid_head}"] = "Value per share (A$): discount rate down, terminal growth across"
    sheet[f"A{grid_head}"].font = BOLD
    sheet[f"{SENS_PV_COLUMN}{grid_head}"] = "PV of cash flows at the rate"
    _note(sheet, grid_head, "The terminal growth rates across the top repeat row 60.")
    for column in columns:
        _put(sheet, f"{column}{grid_head}", f"={column}{R_SENS_GROWTH}", PERCENT)
    terminal_fcf_row = SENS_FIRST_ROW + len(TERMINAL_LABELS) - 1
    fcf_range = _row_range(R_FCF)
    time_range = _row_range(R_TIME)
    for index in range(SENS_SIZE):
        row = SENS_GRID_FIRST_ROW + index
        offset = "=-sensitivity_steps_each_side" if index == 0 else f"=A{row - 1}+1"
        _put(sheet, f"A{row}", offset, "0")
        _put(
            sheet,
            f"B{row}",
            f"=$D${R_RATE}+A{row}*sensitivity_rate_step_pct",
            PERCENT,
        )
        _put(
            sheet,
            f"{SENS_PV_COLUMN}{row}",
            f"=SUMPRODUCT({fcf_range}/(1+$B{row}/100)^{time_range})",
            MONEY,
        )
        for column in columns:
            _put(
                sheet,
                f"{column}{row}",
                f"=(${SENS_PV_COLUMN}{row}"
                f"+{column}${terminal_fcf_row}/(($B{row}-{column}${grid_head})/100)"
                f"*(1+$B{row}/100)^(mid_year_discounting_flag/2)"
                f"/(1+$B{row}/100)^(stub_period_years+${last}${R_YEAR})"
                f"+$D${R_OTHER}-$D${R_NET_DEBT})/$D${R_SHARES}",
                PER_SHARE,
            )
        _note(
            sheet,
            row,
            "A full recalculation at this rate and growth, on the average share count. Column H "
            "holds the present value of the cash flows at the rate; column A, the steps from base.",
        )


def _wacc_block(sheet: Worksheet, first_row: int, risk_free_name: str, stem: str) -> dict[str, str]:
    """The discount rate at the low and high ends of beta and gearing, for one risk-free rate."""
    title = (
        "Discount rate build at my risk-free rate"
        if risk_free_name == "risk_free_rate_pct"
        else "Discount rate build at the expert's risk-free rate"
    )
    _heading(sheet, first_row, title)
    _put(sheet, f"C{first_row}", "Low case")
    _put(sheet, f"D{first_row}", "High case")
    rows = {
        "rf": (first_row + 1, "Risk-free rate", "%", "The 10-year government bond yield."),
        "beta": (
            first_row + 2,
            "Equity beta",
            "number",
            "The low and high ends of the expert's range.",
        ),
        "premium": (first_row + 3, "Market risk premium", "%", "The market risk premium input."),
        "equity": (
            first_row + 4,
            "Cost of equity",
            "%",
            "Risk-free rate plus beta times the premium.",
        ),
        "debt": (
            first_row + 5,
            "Pre-tax cost of debt",
            "%",
            "Risk-free rate plus the debt margin.",
        ),
        "debt_after_tax": (
            first_row + 6,
            "After-tax cost of debt",
            "%",
            "Pre-tax cost less the tax shield.",
        ),
        "weight": (
            first_row + 7,
            "Equity weight",
            "%",
            "Debt takes the rest of the capital structure.",
        ),
        "wacc": (
            first_row + 8,
            "Weighted average cost of capital",
            "%",
            "Weighted cost of equity and after-tax debt.",
        ),
    }
    for row, label, unit, note in rows.values():
        sheet[f"A{row}"] = label
        sheet[f"B{row}"] = unit
        _note(sheet, row, note)
    for column, case in (("C", "low"), ("D", "high")):
        r = {name: f"{column}{row}" for name, (row, *_rest) in rows.items()}
        _put(sheet, r["rf"], f"={risk_free_name}", "0.000")
        _put(sheet, r["beta"], f"=equity_beta_{case}", "0.00")
        _put(sheet, r["premium"], "=market_risk_premium_pct", PERCENT)
        _put(sheet, r["equity"], f"={r['rf']}+{r['beta']}*{r['premium']}", PERCENT)
        _put(sheet, r["debt"], f"={r['rf']}+cost_of_debt_margin_over_risk_free_pct", PERCENT)
        _put(sheet, r["debt_after_tax"], f"={r['debt']}*(1-tax_rate_pct/100)", PERCENT)
        _put(sheet, r["weight"], f"=equity_weight_{case}_pct", PERCENT)
        _put(
            sheet,
            r["wacc"],
            f"={r['weight']}/100*{r['equity']}+(1-{r['weight']}/100)*{r['debt_after_tax']}",
            PERCENT,
        )
    for column_number in range(1, FIRST_YEAR_COLUMN + 2):
        sheet.cell(rows["wacc"][0], column_number).font = BOLD
    wacc_row = rows["wacc"][0]
    return {
        f"{stem}_pct_low": f"{DCF_SHEET}!$C${wacc_row}",
        f"{stem}_pct_high": f"{DCF_SHEET}!$D${wacc_row}",
    }


def add_dcf_sheet(book: Workbook) -> dict[str, str]:
    """Add the DCF sheet and return the cell of each output, keyed by its valuation.json key."""
    sheet: Worksheet = book.create_sheet(DCF_SHEET)
    sheet["A1"] = "Discounted cash flow"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = (
        "Every figure is a formula that reads the Assumptions sheet by name. Change an input "
        "there and this sheet recalculates."
    )
    sheet.column_dimensions["A"].width = LABEL_WIDTH
    sheet.column_dimensions["B"].width = 9
    for column in range(FIRST_YEAR_COLUMN, FIRST_YEAR_COLUMN + YEARS):
        sheet.column_dimensions[get_column_letter(column)].width = 13
    sheet.column_dimensions[NOTE_COLUMN].width = 100
    sheet.freeze_panes = "C5"

    _forecast_block(sheet)

    _heading(
        sheet,
        R_TERMINAL_HEAD,
        "Free cash flow of the year after the forecast (terminal growth rate)",
    )
    _terminal_labels(sheet, TERMINAL_FIRST_ROW)
    terminal_rows = _terminal_year(sheet, "C", TERMINAL_FIRST_ROW, f"$C${R_GROWTH}")
    outputs = _value_block(sheet, f"$C${terminal_rows['fcf']}")
    _sensitivity_block(sheet)
    wacc_outputs = _wacc_block(sheet, R_WACC_HEAD, "risk_free_rate_pct", "wacc")
    wacc_outputs.update(
        _wacc_block(
            sheet,
            R_WACC_HEAD + WACC_BLOCK_HEIGHT,
            "expert_risk_free_rate_pct",
            "wacc_at_expert_risk_free",
        )
    )

    flat: dict[str, str] = dict(wacc_outputs)
    for cells in outputs.values():
        flat.update({key: f"{DCF_SHEET}!${cell[0]}${cell[1:]}" for key, cell in cells.items()})
    flat["dcf_terminal_growth_pct"] = f"{DCF_SHEET}!$D${R_GROWTH}"
    flat["adjusted_net_debt_aud_m"] = f"{DCF_SHEET}!$D${R_NET_DEBT}"
    return flat
