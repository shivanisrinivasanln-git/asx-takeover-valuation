"""The Excel model: an Assumptions sheet filled from the inputs, then sheets of formulas.

Every input cell has an Excel defined name equal to its input name, so a formula reads
`=tax_rate_pct` and not `=Assumptions!C27`. The names of the outputs are the keys of
`valuation.json` (docs/ARCHITECTURE.md). Nothing here calculates a value: the sheets hold
formulas, and a recalculation (scripts/recalculate_with_excel.ps1) fills in the numbers.
"""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

from valuation.forecast import EARLY_YEARS
from valuation.inputs import Inputs
from valuation.workbook_dcf import add_dcf_sheet
from valuation.workbook_multiples import add_multiples_sheet

ASSUMPTIONS_SHEET = "Assumptions"
HEADER_ROW = 4
FIRST_INPUT_ROW = HEADER_ROW + 1
COLUMNS = (
    "Group",
    "Name",
    "Value",
    "Value 2",
    "Value 3",
    "Unit",
    "Source",
    "Location",
    "As at",
    "My judgement",
    "Note",
)
VALUE_COLUMN = 3
SERIES_COLUMNS = (3, 4, 5)
WIDTHS = (22, 52, 14, 10, 10, 18, 40, 28, 16, 14, 90)

GROUP_ASSUMPTIONS = "assumptions.yaml"
GROUP_DATA = "data/extracted"
GROUP_STRUCTURE = "model structure"

INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
BOLD = Font(bold=True)

# Values the engine reads from the transcribed tables, copied here with the row's own source so
# the workbook needs no other file. Names that are also keys of valuation.json use that key.
DATA_INPUTS = (
    ("fy24_revenue_aud_m", "historical_pnl", {"period": "FY24", "line_item": "revenue"}),
    (
        "first_half_fy25_revenue_aud_m",
        "historical_pnl",
        {"period": "1HY25", "line_item": "revenue"},
    ),
    ("first_half_fy25_ebitda_aud_m", "historical_pnl", {"period": "1HY25", "line_item": "ebitda"}),
    (
        "first_half_fy25_depreciation_and_other_amortisation_aud_m",
        "historical_pnl",
        {"period": "1HY25", "line_item": "depreciation_and_other_amortisation"},
    ),
    (
        "expert_business_operations_value_aud_m_low",
        "valuation_summary",
        {"item": "business_operations", "case": "low"},
    ),
    (
        "expert_business_operations_value_aud_m_high",
        "valuation_summary",
        {"item": "business_operations", "case": "high"},
    ),
    (
        "expert_value_per_share_aud_low",
        "valuation_summary",
        {"item": "value_per_share", "case": "low"},
    ),
    (
        "expert_value_per_share_aud_high",
        "valuation_summary",
        {"item": "value_per_share", "case": "high"},
    ),
)


@dataclass(frozen=True)
class InputRow:
    group: str
    name: str
    values: tuple[float, ...]
    unit: str
    source: str
    location: str
    as_at: str
    judgement: bool
    note: str


def input_rows(inputs: Inputs) -> list[InputRow]:
    """Every row of the Assumptions sheet, in the order they are written."""
    rows = [
        InputRow(
            group=GROUP_ASSUMPTIONS,
            name=key,
            values=(entry.value,) if isinstance(entry.value, float) else entry.value,
            unit=entry.unit,
            source=entry.source,
            location=entry.location,
            as_at=entry.as_at,
            judgement=entry.judgement,
            note=entry.note,
        )
        for key, entry in inputs.assumptions.items()
    ]
    for name, table, match in DATA_INPUTS:
        row = inputs.table_row(table, **match)
        rows.append(
            InputRow(
                group=GROUP_DATA,
                name=name,
                values=(float(row["value"]),),
                unit=row["unit"],
                source=row["source"],
                location=row["location"],
                as_at=row["as_at"],
                judgement=False,
                note=f"From data/extracted/{table}.csv. {row['note']}".strip(),
            )
        )
    rows.append(
        InputRow(
            group=GROUP_STRUCTURE,
            name="early_years_count",
            values=(float(EARLY_YEARS),),
            unit="years",
            source="Domain scheme booklet, Grant Samuel report",
            location="p. 142",
            as_at="30 June 2025",
            judgement=False,
            note=(
                "The expert gives the EBITDA margin for FY30 (41%) and FY35 (just below 48%), so "
                "FY26 to FY30 are the five early years and my paths change slope after FY30."
            ),
        )
    )
    return rows


def _define_name(book: Workbook, name: str, reference: str) -> None:
    book.defined_names[name] = DefinedName(name, attr_text=reference)


def add_assumptions_sheet(book: Workbook, inputs: Inputs) -> Worksheet:
    sheet: Worksheet = book.create_sheet(ASSUMPTIONS_SHEET)
    sheet["A1"] = "Assumptions"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = (
        "Every input to the model, with the document, page and date it came from. Shaded cells "
        "are inputs: change one and the other sheets recalculate. A series sits in Value to "
        "Value 3. My own judgements are marked."
    )
    for column, title in enumerate(COLUMNS, start=1):
        cell = sheet.cell(HEADER_ROW, column, title)
        cell.font = BOLD
    for column, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for offset, row in enumerate(input_rows(inputs)):
        number = FIRST_INPUT_ROW + offset
        sheet.cell(number, 1, row.group)
        sheet.cell(number, 2, row.name)
        for column, value in zip(SERIES_COLUMNS, row.values, strict=False):
            cell = sheet.cell(number, column, value)
            cell.fill = INPUT_FILL
        sheet.cell(number, 6, row.unit)
        sheet.cell(number, 7, row.source)
        sheet.cell(number, 8, row.location)
        sheet.cell(number, 9, row.as_at)
        sheet.cell(number, 10, "yes" if row.judgement else "")
        sheet.cell(number, 11, row.note)
        last = SERIES_COLUMNS[len(row.values) - 1]
        first_letter = get_column_letter(VALUE_COLUMN)
        reference = f"'{ASSUMPTIONS_SHEET}'!${first_letter}${number}"
        if last != VALUE_COLUMN:
            reference += f":${get_column_letter(last)}${number}"
        _define_name(book, row.name, reference)
    sheet.freeze_panes = f"C{FIRST_INPUT_ROW}"
    return sheet


def build_workbook(inputs: Inputs) -> Workbook:
    book = Workbook()
    default = book.active
    if default is not None:
        book.remove(default)
    add_assumptions_sheet(book, inputs)
    outputs = add_dcf_sheet(book)
    outputs.update(add_multiples_sheet(book, inputs, "trading"))
    outputs.update(add_multiples_sheet(book, inputs, "precedent"))
    for key, reference in outputs.items():
        _define_name(book, key, reference)
    return book
