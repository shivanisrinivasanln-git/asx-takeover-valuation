"""Every named output cell of the Excel model comes back from Excel's calculation as a number.

The model file holds the values Excel calculated (scripts/build_model.py). These tests read the
stored values and fail on a cell that is missing, empty or an error, or that differs from the
engine's figure by more than the tolerance for its unit. The staleness check, which rebuilds the
workbook from the current inputs and compares it with the committed file, is added with task
P3-15.
"""

from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from valuation.inputs import load_inputs
from valuation.output import build_valuation
from valuation.reconcile import value_problems
from valuation.workbook import build_workbook

ROOT = Path(__file__).resolve().parents[2]
MODEL = ROOT / "model" / "domain_valuation_model.xlsx"


@pytest.fixture(scope="module")
def expected() -> dict[str, dict[str, object]]:
    values: dict[str, dict[str, object]] = build_valuation(load_inputs(ROOT))["values"]
    return values


@pytest.fixture
def calculated() -> Workbook:
    return load_workbook(MODEL, data_only=True)


def cell_of(book: Workbook, key: str) -> tuple[str, str]:
    ((sheet_name, reference),) = book.defined_names[key].destinations
    return sheet_name, reference.replace("$", "")


def test_every_key_has_a_calculated_number_within_tolerance(
    calculated: Workbook, expected: dict[str, dict[str, object]]
) -> None:
    assert len(expected) == 134
    assert value_problems(calculated, expected) == []


def test_an_empty_named_cell_is_reported(
    calculated: Workbook, expected: dict[str, dict[str, object]]
) -> None:
    sheet, reference = cell_of(calculated, "dcf_terminal_value_aud_m_mid")
    calculated[sheet][reference].value = None
    problems = value_problems(calculated, expected)
    assert len(problems) == 1
    assert "is empty" in problems[0]


def test_an_error_cell_is_reported(
    calculated: Workbook, expected: dict[str, dict[str, object]]
) -> None:
    sheet, reference = cell_of(calculated, "wacc_pct_low")
    calculated[sheet][reference].value = "#NAME?"
    problems = value_problems(calculated, expected)
    assert len(problems) == 1
    assert "#NAME?" in problems[0]


def test_a_workbook_that_excel_has_not_calculated_is_reported(
    expected: dict[str, dict[str, object]], tmp_path: Path
) -> None:
    path = tmp_path / "raw.xlsx"
    build_workbook(load_inputs(ROOT)).save(path)
    problems = value_problems(load_workbook(path, data_only=True), expected)
    assert len(problems) > 100
