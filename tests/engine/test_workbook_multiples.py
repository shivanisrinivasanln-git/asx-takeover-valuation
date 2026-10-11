"""The Comparables and Precedents sheets: formulas only, evidence copied with its source."""

from collections import defaultdict
from pathlib import Path
from statistics import median

import pytest
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from valuation.formula_audit import names_used, numeric_constants, sheet_references
from valuation.inputs import Inputs, load_inputs
from valuation.output import build_valuation
from valuation.workbook import build_workbook
from valuation.workbook_multiples import (
    COMPARABLES_SHEET,
    EVIDENCE_COLUMNS,
    NOTE_COLUMN,
    PRECEDENTS_SHEET,
)

ROOT = Path(__file__).resolve().parents[2]
SHEETS = (COMPARABLES_SHEET, PRECEDENTS_SHEET)


@pytest.fixture(scope="module")
def inputs() -> Inputs:
    return load_inputs(ROOT)


@pytest.fixture(scope="module")
def book(inputs: Inputs, tmp_path_factory: pytest.TempPathFactory) -> Workbook:
    path = tmp_path_factory.mktemp("model") / "model.xlsx"
    build_workbook(inputs).save(path)
    return load_workbook(path)


def sheet_of(book: Workbook, name: str) -> Worksheet:
    found = book[name]
    assert isinstance(found, Worksheet)
    return found


def formula_cells(sheet: Worksheet) -> list[tuple[str, str, int]]:
    return [
        (cell.coordinate, cell.value, cell.row)
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]


@pytest.mark.parametrize("name", SHEETS)
def test_the_sheet_has_formulas_and_no_typed_in_numbers(book: Workbook, name: str) -> None:
    cells = formula_cells(sheet_of(book, name))
    assert len(cells) > 50
    typed = {coordinate: numeric_constants(formula) for coordinate, formula, _ in cells}
    assert {coordinate: found for coordinate, found in typed.items() if found} == {}


@pytest.mark.parametrize("name", SHEETS)
def test_every_name_is_defined_and_no_address_into_another_sheet(book: Workbook, name: str) -> None:
    defined = set(book.defined_names)
    for coordinate, formula, _ in formula_cells(sheet_of(book, name)):
        assert names_used(formula) <= defined, coordinate
        assert sheet_references(formula) == [], coordinate


@pytest.mark.parametrize("name", SHEETS)
def test_every_row_with_a_formula_has_a_note(book: Workbook, name: str) -> None:
    sheet = sheet_of(book, name)
    missing = sorted(
        {
            row
            for _, _, row in formula_cells(sheet)
            if not str(sheet[f"{NOTE_COLUMN}{row}"].value or "").strip()
        }
    )
    assert missing == []


def test_the_trading_premium_is_the_input_and_the_precedent_premium_is_zero(
    book: Workbook,
) -> None:
    for measure in ("ebitda", "ebita"):
        trading = book.defined_names[f"trading_{measure}_control_premium_pct"]
        precedent = book.defined_names[f"precedent_{measure}_control_premium_pct"]
        ((trading_sheet, trading_ref),) = trading.destinations
        ((precedent_sheet, precedent_ref),) = precedent.destinations
        assert (
            sheet_of(book, trading_sheet)[trading_ref.replace("$", "")].value
            == "=control_premium_applied_pct"
        )
        assert sheet_of(book, precedent_sheet)[precedent_ref.replace("$", "")].value == 0


def test_every_multiples_output_key_resolves_to_a_cell_on_its_sheet(
    book: Workbook, inputs: Inputs
) -> None:
    keys = [
        key
        for key in build_valuation(inputs)["values"]
        if key.startswith(("trading_", "precedent_"))
    ]
    assert len(keys) == 80
    for key in keys:
        ((sheet_name, _),) = book.defined_names[key].destinations
        expected = COMPARABLES_SHEET if key.startswith("trading_") else PRECEDENTS_SHEET
        assert sheet_name == expected, key


def evidence_rows(sheet: Worksheet) -> dict[str, list[object]]:
    """Rows between the evidence header and the 'Lowest' row, by name."""
    rows: dict[str, list[object]] = {}
    started = False
    for row in sheet.iter_rows(min_col=1, max_col=6):
        label = row[0].value
        if label in ("Company", "Deal"):
            started = True
            continue
        if started:
            if label == "Lowest":
                break
            rows[str(label)] = [cell.value for cell in row[2:6]]
    return rows


def test_the_peer_evidence_equals_the_transcribed_multiples(book: Workbook, inputs: Inputs) -> None:
    written = evidence_rows(sheet_of(book, COMPARABLES_SHEET))
    expected: dict[str, dict[tuple[str, str], float]] = defaultdict(dict)
    for row in inputs.tables["comparable_multiples"]:
        expected[row["company"]][(row["measure"], row["basis"])] = float(row["multiple_x"])
    assert set(written) == set(expected)
    for company, by_column in expected.items():
        for index, column in enumerate(EVIDENCE_COLUMNS):
            assert written[company][index] == by_column.get(column), (company, column)


def test_the_deal_evidence_equals_the_transcribed_multiples(book: Workbook, inputs: Inputs) -> None:
    written = evidence_rows(sheet_of(book, PRECEDENTS_SHEET))
    rows = inputs.tables["precedent_transactions"]
    assert len(written) == len({(row["deal"], row["announced"]) for row in rows})
    seen = 0
    for row in rows:
        name = f"{row['deal']} ({row['announced']})"
        index = EVIDENCE_COLUMNS.index((row["measure"], row["basis"]))
        assert written[name][index] == float(row["multiple_x"]), (name, index)
        seen += 1
    assert seen == len(rows)


def test_the_summary_rows_are_min_median_max_of_each_evidence_column(book: Workbook) -> None:
    sheet = sheet_of(book, COMPARABLES_SHEET)
    labelled = {
        str(row[0].value): [cell.value for cell in row[2:6]]
        for row in sheet.iter_rows(min_col=1, max_col=6)
        if row[0].value in ("Lowest", "Median", "Highest")
    }
    assert set(labelled) == {"Lowest", "Median", "Highest"}
    peers = evidence_rows(sheet)
    first_column = [float(str(values[0])) for values in peers.values()]
    assert median(first_column) > 0
    for function, label in (("MIN", "Lowest"), ("MEDIAN", "Median"), ("MAX", "Highest")):
        assert all(str(cell).startswith(f"={function}(") for cell in labelled[label])
