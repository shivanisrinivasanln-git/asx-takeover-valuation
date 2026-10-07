"""The DCF sheet is formulas only: no typed-in numbers, defined names, and a note on every row."""

from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from valuation.formula_audit import names_used, numeric_constants, sheet_references
from valuation.inputs import load_inputs
from valuation.workbook import build_workbook
from valuation.workbook_dcf import DCF_SHEET, NOTE_COLUMN

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture(scope="module")
def book(tmp_path_factory: pytest.TempPathFactory) -> Workbook:
    path = tmp_path_factory.mktemp("model") / "model.xlsx"
    build_workbook(load_inputs(ROOT)).save(path)
    return load_workbook(path)


def dcf(book: Workbook) -> Worksheet:
    sheet = book[DCF_SHEET]
    assert isinstance(sheet, Worksheet)
    return sheet


def formula_cells(sheet: Worksheet) -> list[tuple[str, str, int]]:
    found = []
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                found.append((cell.coordinate, cell.value, cell.row))
    return found


def test_the_scan_finds_a_typed_in_number_and_allows_the_unit_constants() -> None:
    assert numeric_constants("=C10*1.09") == ["1.09"]
    assert numeric_constants("=tax_rate_pct*7") == ["7"]
    assert numeric_constants("=(C10+D10)/2*(1+rate/100)^0") == []
    assert numeric_constants('="FY"&(INDEX(valuation_date_ymd,1)+1+C5)') == []


def test_the_scan_reads_names_and_sheet_references() -> None:
    assert names_used("=C10*tax_rate_pct/100+$D$5") == {"tax_rate_pct"}
    assert sheet_references("=Assumptions!C5*2") == ["Assumptions!C5"]
    assert sheet_references("=tax_rate_pct*2") == []


def test_the_sheet_has_many_formulas_and_no_typed_in_numbers(book: Workbook) -> None:
    cells = formula_cells(dcf(book))
    assert len(cells) > 300
    typed = {coordinate: numeric_constants(formula) for coordinate, formula, _ in cells}
    assert {coordinate: found for coordinate, found in typed.items() if found} == {}


def test_every_name_a_formula_reads_is_defined(book: Workbook) -> None:
    defined = set(book.defined_names)
    for coordinate, formula, _ in formula_cells(dcf(book)):
        assert names_used(formula) <= defined, (coordinate, names_used(formula) - defined)


def test_formulas_read_inputs_by_name_and_never_by_address(book: Workbook) -> None:
    for coordinate, formula, _ in formula_cells(dcf(book)):
        assert sheet_references(formula) == [], coordinate


def test_every_row_with_a_formula_has_a_note(book: Workbook) -> None:
    sheet = dcf(book)
    without_note = sorted(
        {
            row
            for _, _, row in formula_cells(sheet)
            if not str(sheet[f"{NOTE_COLUMN}{row}"].value or "").strip()
        }
    )
    assert without_note == []


def test_the_sheet_reads_the_inputs_it_needs(book: Workbook) -> None:
    used: set[str] = set()
    for _, formula, _ in formula_cells(dcf(book)):
        used |= names_used(formula)
    for name in (
        "revenue_growth_fy26_to_fy35_pct",
        "tax_rate_pct",
        "terminal_growth_pct",
        "discount_rate_low_pct",
        "discount_rate_high_pct",
        "mid_year_discounting_flag",
        "reported_net_debt_including_leases_aud_m",
        "share_rights_vesting_high_m",
        "sensitivity_rate_step_pct",
    ):
        assert name in used, name
