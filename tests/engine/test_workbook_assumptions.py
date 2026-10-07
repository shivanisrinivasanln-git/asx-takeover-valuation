"""The Assumptions sheet holds every input with its source, and each input has a defined name."""

from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from valuation.inputs import Inputs, load_inputs
from valuation.workbook import (
    ASSUMPTIONS_SHEET,
    DATA_INPUTS,
    FIRST_INPUT_ROW,
    build_workbook,
)

ROOT = Path(__file__).resolve().parents[2]
NAME, UNIT, SOURCE, LOCATION, AS_AT, JUDGEMENT, NOTE = 2, 6, 7, 8, 9, 10, 11


@pytest.fixture(scope="module")
def inputs() -> Inputs:
    return load_inputs(ROOT)


@pytest.fixture(scope="module")
def book(inputs: Inputs, tmp_path_factory: pytest.TempPathFactory) -> Workbook:
    path = tmp_path_factory.mktemp("model") / "model.xlsx"
    build_workbook(inputs).save(path)
    return load_workbook(path)


def sheet(book: Workbook) -> Worksheet:
    found = book[ASSUMPTIONS_SHEET]
    assert isinstance(found, Worksheet)
    return found


def rows_by_name(book: Workbook) -> dict[str, int]:
    ws = sheet(book)
    names = [ws.cell(row, NAME).value for row in range(FIRST_INPUT_ROW, ws.max_row + 1)]
    assert len(names) == len(set(names)), "a name is repeated"
    return {str(name): FIRST_INPUT_ROW + offset for offset, name in enumerate(names)}


def test_the_assumptions_sheet_comes_first(book: Workbook) -> None:
    assert book.sheetnames[0] == ASSUMPTIONS_SHEET


def test_every_input_in_the_file_is_a_row_with_its_value_and_source(
    book: Workbook, inputs: Inputs
) -> None:
    ws = sheet(book)
    rows = rows_by_name(book)
    for key, entry in inputs.assumptions.items():
        row = rows[key]
        values = (entry.value,) if isinstance(entry.value, float) else entry.value
        assert [ws.cell(row, column).value for column in range(3, 3 + len(values))] == list(
            values
        ), key
        assert ws.cell(row, UNIT).value == entry.unit, key
        assert ws.cell(row, SOURCE).value == entry.source, key
        assert ws.cell(row, LOCATION).value == entry.location, key
        assert ws.cell(row, AS_AT).value == entry.as_at, key
        assert ws.cell(row, NOTE).value == entry.note, key
        assert (ws.cell(row, JUDGEMENT).value == "yes") == entry.judgement, key


def test_the_sheet_has_exactly_the_inputs_the_data_rows_and_the_structure_row(
    book: Workbook, inputs: Inputs
) -> None:
    expected = {*inputs.assumptions, *(name for name, _, _ in DATA_INPUTS), "early_years_count"}
    assert set(rows_by_name(book)) == expected


def test_the_data_rows_carry_the_values_and_sources_of_the_tables(
    book: Workbook, inputs: Inputs
) -> None:
    ws = sheet(book)
    rows = rows_by_name(book)
    for name, table, match in DATA_INPUTS:
        source = inputs.table_row(table, **match)
        row = rows[name]
        assert ws.cell(row, 3).value == pytest.approx(float(source["value"])), name
        assert ws.cell(row, SOURCE).value == source["source"], name
        assert ws.cell(row, LOCATION).value == source["location"], name
        assert ws.cell(row, AS_AT).value == source["as_at"], name
    # The 1HY25 depreciation is printed as a negative cost and kept as printed.
    depreciation = rows["first_half_fy25_depreciation_and_other_amortisation_aud_m"]
    assert float(str(ws.cell(depreciation, 3).value)) < 0


def test_every_row_has_a_source_location_and_date(book: Workbook) -> None:
    ws = sheet(book)
    for row in range(FIRST_INPUT_ROW, ws.max_row + 1):
        for column in (SOURCE, LOCATION, AS_AT):
            assert str(ws.cell(row, column).value or "").strip(), (row, column)


def test_each_name_points_at_its_value_cell(book: Workbook) -> None:
    ws = sheet(book)
    rows = rows_by_name(book)
    for name, row in rows.items():
        destinations = list(book.defined_names[name].destinations)
        assert len(destinations) == 1, name
        sheet_name, reference = destinations[0]
        assert sheet_name == ASSUMPTIONS_SHEET
        assert reference.replace("$", "").startswith(f"C{row}"), name
        assert ws.cell(row, 3).value is not None


def test_the_valuation_date_series_is_one_name_over_three_cells(book: Workbook) -> None:
    row = rows_by_name(book)["valuation_date_ymd"]
    ((_, reference),) = book.defined_names["valuation_date_ymd"].destinations
    assert reference.replace("$", "") == f"C{row}:E{row}"
