"""Every valuation.json key is a defined name that resolves to one cell, and no sheet but the
Assumptions sheet holds a typed-in number in a formula."""

from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from valuation.formula_audit import names_used, numeric_constants, sheet_references
from valuation.inputs import Inputs, load_inputs
from valuation.output import FOOTBALL_FIELD_ROWS, PREFERRED_KEY_BY_ROW, build_valuation
from valuation.workbook import ASSUMPTIONS_SHEET, build_workbook, input_rows
from valuation.workbook_dcf import DCF_SHEET
from valuation.workbook_dcf import NOTE_COLUMN as DCF_NOTES
from valuation.workbook_multiples import COMPARABLES_SHEET, PRECEDENTS_SHEET
from valuation.workbook_multiples import NOTE_COLUMN as MULTIPLES_NOTES
from valuation.workbook_summary import FIRST_ROW, SUMMARY_SHEET
from valuation.workbook_summary import NOTE_COLUMN as SUMMARY_NOTES

ROOT = Path(__file__).resolve().parents[2]
NOTE_COLUMNS = {
    DCF_SHEET: DCF_NOTES,
    COMPARABLES_SHEET: MULTIPLES_NOTES,
    PRECEDENTS_SHEET: MULTIPLES_NOTES,
    SUMMARY_SHEET: SUMMARY_NOTES,
}


@pytest.fixture(scope="module")
def inputs() -> Inputs:
    return load_inputs(ROOT)


@pytest.fixture(scope="module")
def book(inputs: Inputs, tmp_path_factory: pytest.TempPathFactory) -> Workbook:
    path = tmp_path_factory.mktemp("model") / "model.xlsx"
    build_workbook(inputs).save(path)
    return load_workbook(path)


def test_the_workbook_has_the_five_sheets_in_order(book: Workbook) -> None:
    assert book.sheetnames == [
        ASSUMPTIONS_SHEET,
        DCF_SHEET,
        COMPARABLES_SHEET,
        PRECEDENTS_SHEET,
        SUMMARY_SHEET,
    ]


def test_every_valuation_key_is_a_name_for_exactly_one_cell(book: Workbook, inputs: Inputs) -> None:
    keys = list(build_valuation(inputs)["values"])
    assert len(keys) == 134
    for key in keys:
        assert key in book.defined_names, key
        destinations = list(book.defined_names[key].destinations)
        assert len(destinations) == 1, key
        sheet_name, reference = destinations[0]
        assert sheet_name in book.sheetnames, key
        assert ":" not in reference, f"{key} names a range"
        assert book[sheet_name][reference.replace("$", "")].value is not None, key


def test_the_defined_names_are_the_inputs_and_the_keys_and_nothing_else(
    book: Workbook, inputs: Inputs
) -> None:
    expected = {row.name for row in input_rows(inputs)} | set(build_valuation(inputs)["values"])
    assert set(book.defined_names) == expected


def test_no_output_key_names_a_cell_on_the_assumptions_sheet_unless_it_is_an_input(
    book: Workbook, inputs: Inputs
) -> None:
    input_names = {row.name for row in input_rows(inputs)}
    for key in build_valuation(inputs)["values"]:
        ((sheet_name, _),) = book.defined_names[key].destinations
        assert (sheet_name == ASSUMPTIONS_SHEET) == (key in input_names), key


def test_no_formula_outside_the_assumptions_sheet_has_a_typed_in_number(book: Workbook) -> None:
    for sheet_name in NOTE_COLUMNS:
        sheet = book[sheet_name]
        assert isinstance(sheet, Worksheet)
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    assert numeric_constants(cell.value) == [], (sheet_name, cell.coordinate)
                    assert names_used(cell.value) <= set(book.defined_names), cell.coordinate
                    assert sheet_references(cell.value) == [], (sheet_name, cell.coordinate)


def test_every_formula_row_has_a_note_on_every_sheet(book: Workbook) -> None:
    for sheet_name, note_column in NOTE_COLUMNS.items():
        sheet = book[sheet_name]
        assert isinstance(sheet, Worksheet)
        rows = {
            cell.row
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        }
        missing = [
            r for r in sorted(rows) if not str(sheet[f"{note_column}{r}"].value or "").strip()
        ]
        assert missing == [], sheet_name


def test_the_summary_rows_are_the_football_field_rows(book: Workbook) -> None:
    sheet = book[SUMMARY_SHEET]
    assert isinstance(sheet, Worksheet)
    for offset, (row_id, label, low_key, high_key) in enumerate(FOOTBALL_FIELD_ROWS):
        row = FIRST_ROW + offset
        assert sheet[f"A{row}"].value == label
        assert sheet[f"C{row}"].value == f"={low_key}"
        assert sheet[f"D{row}"].value == f"={high_key}"
        preferred_key = PREFERRED_KEY_BY_ROW[row_id]
        expected = None if preferred_key is None else f"={preferred_key}"
        assert sheet[f"E{row}"].value == expected, row_id
