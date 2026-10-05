"""The practice workbook matches a direct DCF calculation, and every formula has a note.

Excel itself calculated the values stored in the workbook (scripts/recalculate_with_excel.ps1),
so this test compares Excel's answers with the Python calculation.
"""

from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.formula import DataTableFormula
from openpyxl.worksheet.worksheet import Worksheet

from valuation import finance

WORKBOOK = Path(__file__).resolve().parents[2] / "docs" / "practice" / "one_sheet_dcf.xlsx"
NOTE_COLUMN = 8
TABLE_ROWS = range(27, 32)
TABLE_COLUMNS = range(3, 8)
YEARS = 5
HARBOUR_VALUE_PER_SHARE_AUD = 3.054054


def number(ws: Worksheet, row: int, column: int) -> float:
    value = ws.cell(row, column).value
    assert isinstance(value, int | float)
    return float(value)


def sheet(data_only: bool) -> Worksheet:
    book = load_workbook(WORKBOOK, data_only=data_only)
    ws = book["DCF"]
    assert isinstance(ws, Worksheet)
    return ws


def direct_value_per_share_aud(rate: float, terminal_growth: float) -> float:
    """The same DCF as the sheet, by loop: five growing cash flows, then a Gordon terminal value."""
    first_year_cash_flow, growth, net_debt, shares = 30.0, 0.03, 100.0, 100.0
    cash_flow = first_year_cash_flow
    present_value = 0.0
    for year in range(1, YEARS + 1):
        if year > 1:
            cash_flow *= 1 + growth
        present_value += cash_flow / (1 + rate) ** year
    terminal_value = finance.growing_perpetuity_value_aud_m(
        cash_flow * (1 + terminal_growth), rate * 100, terminal_growth * 100
    )
    enterprise_value = present_value + terminal_value / (1 + rate) ** YEARS
    return finance.equity_value_per_share_aud(enterprise_value, net_debt, shares)


def test_base_case_matches_the_hand_worked_harbour_example() -> None:
    values = sheet(data_only=True)
    assert values["B23"].value == pytest.approx(HARBOUR_VALUE_PER_SHARE_AUD, rel=1e-6)
    assert values["B23"].value == pytest.approx(values["B33"].value, rel=1e-9)


def test_every_cell_of_the_data_table_matches_a_direct_calculation() -> None:
    values = sheet(data_only=True)
    for row in TABLE_ROWS:
        rate = number(values, row, 2)
        for column in TABLE_COLUMNS:
            terminal_growth = number(values, 26, column)
            assert values.cell(row, column).value == pytest.approx(
                direct_value_per_share_aud(rate, terminal_growth), rel=1e-9
            ), (row, column)


def test_the_table_is_an_excel_data_table_over_the_right_input_cells() -> None:
    table = sheet(data_only=False)["C27"].value
    assert isinstance(table, DataTableFormula)
    assert table.ref == "C27:G31"
    assert (table.r1, table.r2) == ("B7", "B6")


def test_the_middle_of_the_table_is_the_base_case() -> None:
    values = sheet(data_only=True)
    assert values["E29"].value == pytest.approx(values["B23"].value, rel=1e-9)
    assert values["B29"].value == pytest.approx(values["B6"].value)
    assert values["E26"].value == pytest.approx(values["B7"].value)


def test_every_formula_has_a_note_in_column_h() -> None:
    ws = sheet(data_only=False)
    without_note = []
    for row in ws.iter_rows():
        for cell in row:
            is_formula = isinstance(cell.value, DataTableFormula) or (
                isinstance(cell.value, str) and cell.value.startswith("=")
            )
            if is_formula and not str(ws.cell(cell.row or 0, NOTE_COLUMN).value or "").strip():
                without_note.append(cell.coordinate)
    assert without_note == []
